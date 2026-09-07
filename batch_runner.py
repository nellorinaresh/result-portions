"""
court_result_batch.py — batch orchestration layer.

Groups extracted results by case type, detects duplicate case numbers,
computes month-range-based output filenames per type, and writes each
type's Result document (Arial 11, one entry after another with a divider,
red-flagged lines for failures/duplicates, no audit line in the final output).
"""

import os
from datetime import date
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from extractor_core import process_file, MONTH_NAMES

FONT_NAME = "Arial"
FONT_SIZE = Pt(11)
RED = RGBColor(0xC0, 0x00, 0x00)


def run_batch(folder, mode, output_dir, workdir):
    """Process every supported file in `folder`, return a dict:
       { case_type: { 'entries': [...], 'flags': [...] } }
    'flags' holds red-flag lines (errors + duplicates), already formatted."""
    supported_ext = ('.doc', '.docx', '.odt', '.pdf')
    files = sorted(
        os.path.join(folder, f) for f in os.listdir(folder)
        if f.lower().endswith(supported_ext)
    )

    groups = {}
    seen_case_numbers = {}  # normalized case_no -> first filename

    def get_group(case_type):
        return groups.setdefault(case_type, {"entries": [], "flags": []})

    for path in files:
        result = process_file(path, mode=mode, workdir=workdir)

        if "error" in result:
            case_type = result.get("case_type", "UNCLASSIFIED")
            g = get_group(case_type)
            g["flags"].append(f'Could not process: {result["file"]} — {result["error"]}. Skipped — please check this file manually.')
            continue

        norm_case_no = result["case_no"].replace(" ", "").upper()
        if norm_case_no in seen_case_numbers:
            g = get_group(result["case_type"])
            g["flags"].append(
                f'Duplicate: {result["case_no"]} already processed from '
                f'{seen_case_numbers[norm_case_no]} — this entry from {result["file"]} skipped.'
            )
            continue
        seen_case_numbers[norm_case_no] = result["file"]

        g = get_group(result["case_type"])
        g["entries"].append(result)

    return groups


def compute_filename(case_type, entries, today):
    dated_entries = [e for e in entries if e.get("date_info")]
    if not dated_entries:
        return f"{case_type} (undated) on {today.day}-{today.month}-{today.year}.docx"

    months = sorted({(e["date_info"]["year"], e["date_info"]["month"]) for e in dated_entries})
    first_y, first_m = months[0]
    last_y, last_m = months[-1]

    if (first_y, first_m) == (last_y, last_m):
        range_str = f"{MONTH_NAMES[first_m]}, {first_y}"
    else:
        range_str = f"{MONTH_NAMES[first_m]}, {first_y} to {MONTH_NAMES[last_m]}, {last_y}"

    return f"{case_type} {range_str} on {today.day}-{today.month}-{today.year}.docx"


def _field_para(doc, label, value):
    p = doc.add_paragraph()
    r1 = p.add_run(f"{label}: ")
    r1.bold = True
    r1.font.name = FONT_NAME
    r1.font.size = FONT_SIZE
    r2 = p.add_run(value)
    r2.font.name = FONT_NAME
    r2.font.size = FONT_SIZE
    return p


def _content_label(doc):
    p = doc.add_paragraph()
    r = p.add_run("Content:")
    r.bold = True
    r.font.name = FONT_NAME
    r.font.size = FONT_SIZE


def _content_lines(doc, lines):
    for line in lines:
        p = doc.add_paragraph()
        r = p.add_run(line)
        r.font.name = FONT_NAME
        r.font.size = FONT_SIZE


def _separator(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(10)
    pPr = p._p.get_or_add_pPr()
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '999999')
    pBdr.append(bottom)
    pPr.append(pBdr)


def _flag_para(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.name = FONT_NAME
    r.font.size = FONT_SIZE
    r.font.color.rgb = RED


def build_output_docx(case_type, group, output_dir, today):
    entries = group["entries"]
    flags = group["flags"]

    def sort_key(e):
        import re
        m = re.search(r'(\d+)\s*(?:/|of)\s*(\d{2,4})', e["case_no"])
        num, yr = int(m.group(1)), int(m.group(2))
        if yr < 100:
            yr += 2000
        return (yr, num)

    entries_sorted = sorted(entries, key=sort_key)

    doc = Document()
    style = doc.styles['Normal']
    style.font.name = FONT_NAME
    style.font.size = FONT_SIZE

    title = doc.add_paragraph()
    tr = title.add_run(f"RESULT — {case_type}")
    tr.bold = True
    tr.font.name = FONT_NAME
    tr.font.size = Pt(14)

    for e in entries_sorted:
        pet_label = "Petitioners" if len(e["petitioners"]) > 1 else "Petitioner"
        resp_label = "Respondents" if len(e["respondents"]) > 1 else "Respondent"
        if case_type == "OS":
            pet_label = "Plaintiffs" if len(e["petitioners"]) > 1 else "Plaintiff"
            resp_label = "Defendants" if len(e["respondents"]) > 1 else "Defendant"

        _field_para(doc, "Case No.", e["case_no"])
        _field_para(doc, pet_label, ", ".join(e["petitioners"]))
        _field_para(doc, resp_label, ", ".join(e["respondents"]))
        _field_para(doc, "Date of Disposal", e["date_info"]["display"] if e["date_info"] else "Unknown")
        _content_label(doc)
        _content_lines(doc, e["content"])
        _separator(doc)

    for flag_text in flags:
        _flag_para(doc, flag_text)

    fname = compute_filename(case_type, entries_sorted, today)
    out_path = os.path.join(output_dir, fname)
    doc.save(out_path)
    return out_path
